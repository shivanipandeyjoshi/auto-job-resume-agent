from __future__ import annotations

from pathlib import Path

from docx import Document

from models.job_models import ApplicationState, Metadata, OptimizedResume
from utils.config_manager import ConfigManager
from utils.csv_manager import CsvManager
from utils.date_utils import today_str
from utils.docx_service import DocxService
from utils.logger import AppLogger, resolve_logger
from utils.pdf_service import PdfService
from utils.text_utils import slugify


class GenerateResumeAgent:
    """Generate DOCX and PDF outputs from the optimized resume content."""

    def __init__(self, config_manager: ConfigManager | None = None, logger: AppLogger | None = None) -> None:
        self.config_manager = config_manager or ConfigManager()
        self.logger = resolve_logger(logger, self.config_manager)
        self.generated_dir = Path(self.config_manager.get("project.generated_dir", "resumes/generated"))
        self.generated_dir.mkdir(parents=True, exist_ok=True)
        self.docx_service = DocxService(str(self.config_manager.get("project.resume_template", "templates/master_resume.docx")))
        self.pdf_service = PdfService()

    def generate(self, state: ApplicationState) -> ApplicationState:
        """Create a DOCX document from the optimized resume and record metadata."""

        resume = state.optimized_resume or OptimizedResume()
        if not resume.summary and state.resume:
            resume.summary = state.resume.summary
        if not resume.summary:
            resume.summary = "Experienced QA and testing professional with strong skills in automation, regression, and cross-functional delivery."
        if not resume.skills and state.resume:
            resume.skills = state.resume.skills
        if not resume.skills:
            resume.skills = ["Testing", "QA", "Automation", "Python", "Selenium"]
        if not resume.experience and state.resume:
            resume.experience = state.resume.experience
        if not resume.experience:
            resume.experience = []
        if not resume.projects and state.resume:
            resume.projects = state.resume.projects
        latest_job = state.ranked_jobs[0]["job"] if state.ranked_jobs else None
        position = latest_job.title if latest_job else (state.resume.summary.split(".")[0] if state.resume and state.resume.summary else "Experienced Professional")
        company = latest_job.company if latest_job else "Target Company"
        stamp = today_str()
        safe_position = slugify(position)
        safe_company = slugify(company)

        docx_path = self.generated_dir / f"{safe_position}_{safe_company}_{stamp}.docx"
        pdf_path = self.generated_dir / f"{safe_position}_{safe_company}_{stamp}.pdf"

        template_path = Path(self.config_manager.get("project.resume_template", "templates/master_resume.docx"))
        if not template_path.exists():
            template_path = self._create_template(template_path)

        document = Document(template_path)
        replacements = {
            "{{name}}": position,
            "{{company}}": company,
            "{{summary}}": resume.summary,
            "{{skills}}": ", ".join(resume.skills),
            "{{experience}}": "\n".join(
                f"- {entry.role} at {entry.company}: {', '.join(entry.highlights)}" for entry in resume.experience
            ),
            "{{projects}}": "\n".join(
                f"- {project.name}: {project.description}" for project in resume.projects
            ),
        }
        self.docx_service.replace_text(document, replacements)
        document.save(docx_path)
        self.pdf_service.convert(docx_path, pdf_path)
        self.logger.info("Generated resume files at {} and {}", docx_path, pdf_path)

        metadata = Metadata(
            date=stamp,
            company=company,
            position=position,
            resume_path=str(docx_path),
            pdf_path=str(pdf_path),
            job_url=latest_job.source_url if latest_job else "",
            apply_url="",
            ranking=state.ranked_jobs[0]["score"] if state.ranked_jobs else 0.0,
            ats_score=resume.ats_score,
            status="generated",
            notes="Generated locally",
        )
        state.metadata.append(metadata)
        state.generated_files = [str(docx_path), str(pdf_path)]

        return state

    def _create_template(self, template_path: Path) -> Path:
        """Create a sample DOCX template when no template file exists."""

        template_path.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        document.add_heading("{{name}}", level=1)
        document.add_paragraph("{{company}}")
        document.add_heading("Summary", level=2)
        document.add_paragraph("{{summary}}")
        document.add_heading("Skills", level=2)
        document.add_paragraph("{{skills}}")
        document.add_heading("Experience", level=2)
        document.add_paragraph("{{experience}}")
        document.add_heading("Projects", level=2)
        document.add_paragraph("{{projects}}")
        document.save(template_path)
        return template_path
