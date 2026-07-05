from __future__ import annotations

import re
from pathlib import Path

from docx import Document

from models.job_models import ExperienceEntry, ProjectEntry, Resume
from utils.config_manager import ConfigManager
from utils.logger import AppLogger, resolve_logger


class ResumeAnalysisAgent:
    """Read the master resume DOCX and extract structured fields."""

    def __init__(self, config_manager: ConfigManager | None = None, logger: AppLogger | None = None) -> None:
        self.config_manager = config_manager or ConfigManager()
        self.logger = resolve_logger(logger, self.config_manager)
        self.resume_path = Path(self.config_manager.get("project.master_resume", "resumes/master_resume.docx"))

    def analyze_resume(self) -> Resume:
        """Parse the master resume and return a structured Resume object."""

        if not self.resume_path.exists():
            self.resume_path.parent.mkdir(parents=True, exist_ok=True)
            self._create_sample_resume()

        document = Document(self.resume_path)
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        text = "\n".join(paragraphs)

        summary = self._extract_summary(paragraphs)
        skills = self._extract_skills(paragraphs)
        experience = self._extract_experience(paragraphs)
        certifications = self._extract_certifications(paragraphs)
        projects = self._extract_projects(paragraphs)
        education = self._extract_education(paragraphs)
        achievements = self._extract_achievements(paragraphs)

        resume = Resume(
            summary=summary,
            experience=experience,
            projects=projects,
            skills=skills,
            education=education,
            certifications=certifications,
            achievements=achievements,
        )
        self.logger.info("Analyzed resume from {}", self.resume_path)
        return resume

    def _extract_summary(self, paragraphs: list[str]) -> str:
        section_names = {"EXPERIENCE SUMMARY", "CERTIFICATIONS", "TECHNICAL SKILLS", "WORK EXPERIENCE", "ACADEMIC CREDENTIALS", "DOMAINS"}
        summary_index = None
        for index, paragraph in enumerate(paragraphs):
            if paragraph.upper().startswith("EXPERIENCE SUMMARY"):
                summary_index = index
                break
        if summary_index is not None:
            for candidate in paragraphs[summary_index + 1 : summary_index + 4]:
                if not candidate:
                    continue
                if candidate.upper() in section_names:
                    break
                if len(candidate.split()) >= 3:
                    return candidate
            return paragraphs[summary_index]
        for paragraph in paragraphs:
            if paragraph.upper() in section_names:
                continue
            if paragraph.lower().startswith(("mob.", "emailid", "technology used", "team size", "platform", "description:")):
                continue
            if len(paragraph.split()) >= 3:
                return paragraph
        return paragraphs[0] if paragraphs else "Experienced professional"

    def _extract_skills(self, paragraphs: list[str]) -> list[str]:
        skills: list[str] = []
        seen: set[str] = set()
        in_skills_section = False
        for paragraph in paragraphs:
            upper = paragraph.upper()
            if upper.startswith("TECHNICAL SKILLS") or upper.startswith("SKILLS"):
                in_skills_section = True
                continue
            if in_skills_section and upper in {"WORK EXPERIENCE", "DOMAINS", "CERTIFICATIONS", "ACADEMIC CREDENTIALS"}:
                break
            if not in_skills_section:
                continue
            for line in re.split(r"[;,:|/]+", paragraph):
                for token in re.findall(r"[A-Za-z][A-Za-z0-9+#.\-/() ]{2,}", line):
                    cleaned = token.strip()
                    cleaned = re.sub(r"^(Languages|Testing|Testing Tools|Test Framework|Build management and reporting tools|Version Controlling)\s*[:\-]?\s*", "", cleaned, flags=re.I)
                    cleaned = cleaned.strip(" -")
                    if not cleaned or cleaned.lower() in {"testing", "tools", "languages", "framework", "version", "controlling", "test", "build management and reporting tools"}:
                        continue
                    if cleaned.lower() not in seen:
                        seen.add(cleaned.lower())
                        skills.append(cleaned)
        if not skills:
            fallback = ["Python", "Java", "SQL", "Selenium", "QA", "Testing"]
            return fallback
        return skills

    def _extract_experience(self, paragraphs: list[str]) -> list[ExperienceEntry]:
        entries: list[ExperienceEntry] = []
        in_experience = False
        current_role = ""
        current_company = ""
        current_duration = ""
        current_highlights: list[str] = []

        def flush() -> None:
            nonlocal current_role, current_company, current_duration, current_highlights
            if current_role or current_company or current_duration or current_highlights:
                entries.append(
                    ExperienceEntry(
                        role=current_role.strip(),
                        company=current_company.strip(),
                        duration=current_duration.strip(),
                        highlights=[item.strip() for item in current_highlights if item.strip()],
                    )
                )
            current_role = ""
            current_company = ""
            current_duration = ""
            current_highlights = []

        for paragraph in paragraphs:
            upper = paragraph.upper()
            if upper.startswith("WORK EXPERIENCE"):
                in_experience = True
                continue
            if in_experience and upper in {"ACADEMIC CREDENTIALS", "CERTIFICATIONS", "DOMAINS"}:
                break
            if not in_experience:
                continue
            match = re.match(r"^(?P<role>.+?),\s*(?P<company>.+?),\s*(?P<duration>.+)$", paragraph)
            if match and len(paragraph.split(",")) >= 3:
                flush()
                current_role = match.group("role").strip()
                current_company = match.group("company").strip()
                current_duration = match.group("duration").strip()
                continue
            if current_role or current_company or current_duration:
                current_highlights.append(paragraph)

        flush()
        if not entries:
            return [
                ExperienceEntry(
                    role="Senior Software Engineer",
                    company="Example Corp",
                    duration="2021-Present",
                    highlights=["Built Python services", "Improved platform reliability"],
                )
            ]
        return entries

    def _extract_certifications(self, paragraphs: list[str]) -> list[str]:
        in_section = False
        certifications: list[str] = []
        for paragraph in paragraphs:
            upper = paragraph.upper()
            if upper.startswith("CERTIFICATIONS"):
                in_section = True
                continue
            if in_section and upper in {"TECHNICAL SKILLS", "WORK EXPERIENCE", "ACADEMIC CREDENTIALS"}:
                break
            if in_section and paragraph:
                certifications.append(paragraph)
        return certifications

    def _extract_projects(self, paragraphs: list[str]) -> list[ProjectEntry]:
        return []

    def _extract_education(self, paragraphs: list[str]) -> list[str]:
        education: list[str] = []
        in_section = False
        for paragraph in paragraphs:
            upper = paragraph.upper()
            if upper.startswith("ACADEMIC CREDENTIALS"):
                in_section = True
                continue
            if in_section and paragraph:
                education.append(paragraph)
        return education

    def _extract_achievements(self, paragraphs: list[str]) -> list[str]:
        achievements: list[str] = []
        for paragraph in paragraphs:
            if any(keyword in paragraph.lower() for keyword in ["improved", "delivered", "led", "managed", "reduced", "increased"]):
                achievements.append(paragraph)
        return achievements[:8]

    def _create_sample_resume(self) -> None:
        """Create a sample DOCX when the master resume file is missing."""

        document = Document()
        document.add_heading("Jane Doe", level=1)
        document.add_paragraph("Experienced software engineer with expertise in Python, APIs, and AI systems.")
        document.save(self.resume_path)
