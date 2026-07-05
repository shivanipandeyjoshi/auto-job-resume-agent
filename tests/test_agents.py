from pathlib import Path

import yaml
from docx import Document

from agents.discover_jobs import DiscoverJobsAgent
from agents.resume_analysis import ResumeAnalysisAgent
from utils.config_manager import ConfigManager
from utils.logger import AppLogger


def test_resume_analysis_agent_reads_master_resume():
    manager = ConfigManager()
    agent = ResumeAnalysisAgent(config_manager=manager)
    resume = agent.analyze_resume()
    assert resume.summary
    assert resume.skills


def test_discover_jobs_uses_configured_locations_and_job_types(tmp_path, monkeypatch):
    config_path = tmp_path / "config.yaml"
    config_path.write_text(
        yaml.safe_dump(
            {
                "project": {"root_dir": ".", "name": "test"},
                "agents": {
                    "discover_jobs": {
                        "keywords": ["QA Engineer", "Software Tester"],
                        "locations": ["India", "Remote"],
                        "portals": ["LinkedIn"],
                        "job_types": ["Hybrid", "Full-time"],
                        "limit": 2,
                    }
                },
            }
        ),
        encoding="utf-8",
    )
    monkeypatch.chdir(tmp_path)

    manager = ConfigManager(config_path=str(config_path))
    agent = DiscoverJobsAgent(config_manager=manager)
    state = agent.discover_jobs(None)

    assert state.job_summaries
    assert state.job_summaries[0].location == "India"
    assert state.job_summaries[0].employment_type == "Hybrid"


def test_build_search_queries_uses_resume_keywords():
    manager = ConfigManager()
    agent = DiscoverJobsAgent(config_manager=manager)
    resume_text = "Senior QA Engineer with 10+ years in Selenium, Python, banking and telecom testing"

    queries = agent._build_search_queries(resume_text)

    assert any("qa engineer" in query.lower() for query in queries)
    assert any("selenium" in query.lower() for query in queries)
    assert any("banking" in query.lower() for query in queries)
    assert any("telecom" in query.lower() for query in queries)


def test_build_search_queries_include_semantic_terms_from_resume():
    manager = ConfigManager()
    agent = DiscoverJobsAgent(config_manager=manager)
    resume_text = "Senior QA Automation Engineer with Playwright, Python, API testing, banking and telecom"

    queries = agent._build_search_queries(resume_text)

    assert any("playwright" in query.lower() for query in queries)
    assert any("api" in query.lower() for query in queries)
    assert any("automation" in query.lower() for query in queries)


def test_app_logger_can_be_passed_to_agents():
    manager = ConfigManager()
    logger = AppLogger(manager)
    agent = ResumeAnalysisAgent(config_manager=manager, logger=logger)

    assert agent.logger is logger.get_logger()


def test_resume_analysis_extracts_real_skills_and_experience(tmp_path):
    resume_path = tmp_path / "resume.docx"
    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("Senior QA Engineer")
    doc.add_paragraph("EXPERIENCE SUMMARY")
    doc.add_paragraph("10+ years in software testing and automation")
    doc.add_paragraph("TECHNICAL SKILLS")
    doc.add_paragraph("Languages: Python, Java")
    doc.add_paragraph("Testing: Selenium, Playwright")
    doc.add_paragraph("WORK EXPERIENCE")
    doc.add_paragraph("QA Lead, Example Corp, Jan 2020 – Present")
    doc.add_paragraph("Led automation and regression testing")
    doc.add_paragraph("Improved release quality")
    doc.save(resume_path)

    config_path = tmp_path / "config.yaml"
    config_path.write_text("project:\n  master_resume: \"{}\"\n".format(resume_path), encoding="utf-8")
    manager = ConfigManager(config_path=str(config_path))
    agent = ResumeAnalysisAgent(config_manager=manager)
    resume = agent.analyze_resume()

    assert resume.summary.startswith("10+ years")
    assert any(skill.lower() == "python" for skill in resume.skills)
    assert any(skill.lower() == "selenium" for skill in resume.skills)
    assert resume.experience[0].role == "QA Lead"
    assert resume.experience[0].highlights[0].startswith("Led")
